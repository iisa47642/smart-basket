from pathlib import Path
from docx import Document

path = Path(__file__).resolve().parents[1] / "report_input.docx"
doc = Document(path)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print("STYLES")
for style in doc.styles:
    if style.type == 1 and style.name in {"Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"}:
        font = style.font
        print(style.name, font.name, font.size.pt if font.size else None, font.bold, font.italic)

print("HEADINGS")
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style else ""
    if text and (style.startswith("Heading") or style in {"Title", "Subtitle"}):
        print(f"{i:04d}\t{style}\t{text}")

print("LAST_PARAGRAPHS")
for i, paragraph in list(enumerate(doc.paragraphs))[-80:]:
    text = paragraph.text.strip()
    if text:
        print(f"{i:04d}\t{paragraph.style.name}\t{text}")

print("SECTIONS")
for i, section in enumerate(doc.sections):
    print(
        i,
        "size", section.page_width, section.page_height,
        "margins", section.top_margin, section.right_margin, section.bottom_margin, section.left_margin,
        "start", section.start_type,
    )
